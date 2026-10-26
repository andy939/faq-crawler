#!/usr/bin/env python3
"""
本機工作台 —— 把 docs/ 這個網頁跑起來，並且讓網頁上的按鈕可以真的去抓資料。

    python 工作台.py          # 開 http://127.0.0.1:8765，只有這台電腦能用
    python 工作台.py --lan    # 家裡區網的其他電腦也能連（區網以外一律擋掉）

跟 GitHub Pages 上那頁是同一個 index.html。差別只在：
放在 GitHub 上是唯讀的監控頁（靜態網頁沒有後端，按不動），
在這裡打開就會多出整塊控制台 —— 選抓法、筆數、日期、限速，按下去抓，
八個指標即時跳動，可以中途停止。

抓取本身是開一個子行程去跑 crawl.py，所以行為跟你雙擊 更新.bat
或 GitHub Actions 上跑的完全一樣，不會有第二套邏輯。
"""

import argparse
import ipaddress
import json
import os
import socket
import subprocess
import sys
import threading
import time
import urllib.parse
import webbrowser
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer

import crawl as C

HERE = os.path.dirname(os.path.abspath(__file__))
DOCS = os.path.join(HERE, "docs")
VEC = os.path.join(HERE, "vectors")
PORT = 8765

# 檢索索引常駐在記憶體裡。BM25 索引建一次要一秒多、語意模型要載入好幾秒，
# 每次查詢重建的話打一個字就卡住。第一次用到才建（import 檢索 會把 numpy
# 拉進來，放在啟動時做會讓工作台變慢）。
_IX = None
_IX_LOCK = threading.Lock()


def index():
    global _IX
    with _IX_LOCK:
        if _IX is None:
            import 檢索
            _IX = 檢索.Index()
        return _IX


def warm_up():
    """背景把索引和第一個語意模型載好。不做的話，使用者第一次查詢要等
    三十幾秒（載模型），之後才變成幾十毫秒 —— 那三十秒全砸在第一個人身上。
    開在背景執行緒，工作台本身照常馬上可以用。"""
    def go():
        try:
            ix = index()
            ix.bm25("預熱", k=1)
            tags = vector_tags()
            if tags:
                t0 = time.perf_counter()
                ix.semantic("預熱", tags[0]["tag"], k=1)
                print(f"語意查詢就緒：{tags[0]['tag']}"
                      f"（{time.perf_counter() - t0:.0f} 秒）")
        except Exception as e:
            print(f"語意查詢無法使用：{type(e).__name__}: {e}")
    threading.Thread(target=go, daemon=True).start()


def vector_tags():
    """算好的向量有哪些。沒有 vectors/ 資料夾就是還沒算過。"""
    if not os.path.isdir(VEC):
        return []
    out = []
    for f in sorted(os.listdir(VEC)):
        # .state.json 是逐筆記錄（增量重算與放寬保護用），不是一組向量的說明檔
        if f.endswith(".json") and not f.endswith(".state.json"):
            try:
                out.append(json.load(open(os.path.join(VEC, f),
                                          encoding="utf-8")))
            except Exception:
                pass
    return out

# 網頁上那排工具按鈕。key 是網頁送來的名字，值是要跑的命令列。
TOOLS = {
    "check":  ("環境檢查", [sys.executable, "-u", "crawl.py", "--check"]),
    "fill":   ("補齊（官網幾筆就抓幾筆）", [sys.executable, "-u", "crawl.py"]),
    "full":   ("全站重抓", [sys.executable, "-u", "crawl.py", "--full"]),
    "report": ("產生分析報表", [sys.executable, "-u", "report.py"]),
    "excel":  ("產生 Excel", [sys.executable, "-u", "to_excel.py"]),
    "sync":   ("同步到網站", [sys.executable, "-u", "同步網站.py"]),
}

# 抓完自動接著跑的後續工作（網頁上那個勾選框）
FOLLOW = {"sync": ("同步到網站", [sys.executable, "-u", "同步網站.py"])}


class Job:
    """同時只跑一個。抓取中再按一次不會疊上去，只會被擋掉。"""

    def __init__(self):
        self.lock = threading.Lock()
        self.lines = []
        self.metrics = {}
        self.proc = None
        self.label = ""
        self.cmd = ""
        self.started = 0.0
        self.code = None
        self.then = None          # 跑完要不要接著做別的（例如同步到網站）

    @property
    def running(self):
        return self.proc is not None and self.proc.poll() is None

    def start(self, label, args, then=None):
        with self.lock:
            if self.running:
                return False, "已經有一輪在跑了"
            if os.path.exists(C.STOP):
                os.remove(C.STOP)
            self.cmd = " ".join(a for a in args[2:]) or "（預設）"
            self.lines = [f"$ python {' '.join(args[2:])}".rstrip()]
            self.metrics = {}
            self.label, self.started, self.code = label, time.time(), None
            self.then = then
            env = dict(os.environ, CRAWL_JSON="1", PYTHONIOENCODING="utf-8")
            self.proc = subprocess.Popen(
                args, cwd=HERE, stdout=subprocess.PIPE,
                stderr=subprocess.STDOUT, text=True, encoding="utf-8",
                errors="replace", bufsize=1, env=env,
                # Windows 下不要另外彈出一個主控台視窗
                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        threading.Thread(target=self._pump, daemon=True).start()
        return True, label

    def stop(self):
        """建一個旗標檔，crawl.py 每次請求前會看。這樣是「好好地停」——
        已經抓到的會存檔，不像直接砍掉行程什麼都不剩。"""
        if not self.running:
            return False
        open(C.STOP, "w").close()
        self.lines.append("—— 已要求停止，等這一筆結束後收尾… ——")
        return True

    def _pump(self):
        for raw in self.proc.stdout:
            if raw.startswith("@@"):          # 機器讀的即時指標
                try:
                    self.metrics.update(json.loads(raw[2:]))
                except ValueError:
                    pass
                continue
            # crawl.py 的進度是用 \r 原地更新的，這裡拆成一行一行
            for line in raw.replace("\r", "\n").split("\n"):
                if line.strip():
                    self.lines.append(line.rstrip())
            del self.lines[:-600]             # 只留最後 600 行
        self.code = self.proc.wait()
        el = time.time() - self.started
        self.lines.append(
            f"—— 結束（{'成功' if self.code == 0 else f'離開碼 {self.code}'}），"
            f"共 {el:.0f} 秒 ——")
        # 抓完自動接著同步（網頁勾了「抓完自動同步」才會有）。
        # 只有順利跑完才接 —— 被中斷或挨罰跳過的那輪，資料不完整，
        # 不該把它推到公開網站上。
        if self.then and self.code == 0:
            nxt, self.then = self.then, None
            prev = self.lines
            self.lines = prev + ["", "—— 接著自動同步到網站 ——"]
            keep = self.lines
            self.start(*nxt)
            self.lines = keep + self.lines[1:]
        elif self.then:
            self.then = None
            self.lines.append("（這輪沒有順利跑完，不自動同步）")

    def state(self, since=0):
        return {
            "running": self.running, "label": self.label, "cmd": self.cmd,
            "code": self.code, "metrics": self.metrics,
            "elapsed": round(time.time() - self.started, 1) if self.started else 0,
            "total": len(self.lines), "lines": self.lines[since:],
        }


JOB = Job()


def build_args(q):
    """把網頁送來的參數組成 crawl.py 的命令列。"""
    args = [sys.executable, "-u", "crawl.py"]
    mode = q.get("mode", ["fill"])[0]
    if mode == "full":
        args.append("--full")
    elif mode == "refresh":
        args += ["--refresh", str(int(q.get("refresh", ["300"])[0] or 300))]
    elif mode == "check":
        return "環境檢查", args + ["--check"]
    elif mode == "dry":
        args.append("--dry-run")
    elif mode == "rss":
        args.append("--rss")

    method = q.get("method", ["B"])[0]
    if method in C.METHODS:
        args += ["--method", method]
    limit = q.get("limit", [""])[0]
    if limit and limit.isdigit() and int(limit) > 0:
        args += ["--limit", limit]
    for k, flag in (("from", "--from"), ("to", "--to")):
        v = (q.get(k, [""])[0] or "").strip()
        if v:
            args += [flag, v]
    rate = q.get("rate", [""])[0]
    if rate and rate.isdigit():
        args += ["--rate-ms", rate]
    if q.get("force", [""])[0] == "1":
        args.append("--force")

    label = {"fill": "補齊", "full": "全站重抓", "refresh": "補齊＋複查",
             "dry": "試跑", "rss": "RSS 快更"}.get(mode, mode)
    label += f"・抓法 {method}"
    if limit:
        label += f"・{limit} 筆"
    return label, args


def write_xlsx(body):
    """把網頁送來的列寫成 Excel，回傳檔名。

    body = {name, sheet, cols:[欄位名], rows:[[值, …], …]}
    欄寬照內容自動估，第一列凍結並開篩選 —— 開起來就能直接用。
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    cols = body.get("cols") or []
    rows = body.get("rows") or []
    wb = Workbook()
    ws = wb.active
    ws.title = (body.get("sheet") or "資料")[:31]
    ws.append(cols)
    ctrl = __import__("re").compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
    for r in rows:
        ws.append([ctrl.sub("", v)[:32000] if isinstance(v, str) else v
                   for v in r])
    fill = PatternFill("solid", fgColor="1A6A5A")
    font = Font(bold=True, color="FFFFFF")
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=i)
        cell.fill, cell.font = fill, font
        cell.alignment = Alignment(vertical="center")
        w = max([len(str(c)) * 2] + [len(str(r[i - 1])) for r in rows[:200]
                                     if i <= len(r)][:200] or [10])
        ws.column_dimensions[get_column_letter(i)].width = min(max(w, 10), 60)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions

    base = (body.get("name") or "匯出").replace("/", "_").replace("\\", "_")
    name = f"{base}_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb.save(os.path.join(HERE, name))
    return name


def latest_wire():
    """從 exports/ 的逐筆測量檔撈出每一筆的「抓下 byte」，新的蓋舊的。"""
    out = {}
    try:
        import csv
        import glob
        for p in sorted(glob.glob(os.path.join(HERE, "exports", "*.csv"))):
            with open(p, encoding="utf-8-sig", newline="") as f:
                for r in csv.DictReader(f):
                    if r.get("sid") and r.get("抓下bytes"):
                        out[r["sid"]] = int(float(r["抓下bytes"]))
    except Exception:
        pass
    return out


class Handler(SimpleHTTPRequestHandler):
    def __init__(self, *a, **k):
        super().__init__(*a, directory=DOCS, **k)

    def _json(self, obj, code=200):
        body = json.dumps(obj, ensure_ascii=False).encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(body)

    def do_GET(self):
        path, _, qs = self.path.partition("?")
        q = urllib.parse.parse_qs(qs)
        if path == "/api/ping":
            # 網頁靠這支判斷「我是不是跑在本機工作台上」，
            # 在 GitHub Pages 上這支會 404，網頁就只顯示唯讀的監控。
            return self._json({
                "local": True,
                "methods": [{"key": k, "label": v[0], "desc": v[1],
                             "workers": v[2].get("workers", 1)}
                            for k, v in C.METHODS.items()],
                "tools": [{"key": k, "label": v[0]} for k, v in TOOLS.items()],
                "rate_default": C.RATE_MS,
            })
        if path == "/api/log":
            since = int(q.get("since", ["0"])[0] or 0)
            return self._json(JOB.state(since))
        if path == "/api/download":
            # 讓網頁把產生的 Excel／報表直接下載回去，不用自己去翻資料夾。
            # 只給專案根目錄下的 .xlsx / .md，而且不接受路徑分隔符。
            name = os.path.basename(q.get("name", [""])[0])
            full = os.path.join(HERE, name)
            if (not name or os.sep in name or not os.path.isfile(full)
                    or os.path.splitext(name)[1].lower() not in (".xlsx", ".md")):
                return self.send_error(404)
            with open(full, "rb") as f:
                body = f.read()
            self.send_response(200)
            self.send_header("Content-Type", "application/octet-stream")
            self.send_header("Content-Length", str(len(body)))
            self.send_header("Content-Disposition",
                             "attachment; filename*=UTF-8''"
                             + urllib.parse.quote(name))
            self.end_headers()
            self.wfile.write(body)
            return
        if path == "/api/xlsx":
            return self._json({"ok": False, "msg": "請用 POST"}, 405)
        if path == "/api/samples":
            # 逐筆的「抓下 byte」只存在 exports/ 的測量檔裡（那是每次抓取的
            # 量測值，不是問答本身的屬性，所以沒進 faq.json）。
            # 網頁那張逐筆表要顯示它，就從最近的測量檔撈出來，新的蓋舊的。
            return self._json({"wire": latest_wire()})
        if path == "/api/search":
            # 語意查詢。只排序、不生成答案 —— 回去的是 faq.json 裡的原文那幾筆。
            # 沒帶 q 就只回「有哪些模型可以選」，網頁用它決定要不要顯示那個選單。
            tags = vector_tags()
            kw = (q.get("q") or [""])[0].strip()
            if not kw:
                return self._json({"ok": True, "vectors": tags})
            mode = (q.get("mode") or ["bm25"])[0]
            tag = (q.get("tag") or [""])[0] or (tags[0]["tag"] if tags else "")
            k = max(1, min(int((q.get("k") or ["50"])[0] or 50), 200))
            t0 = time.perf_counter()
            try:
                rows = index().search(kw, mode=mode, tag=tag, k=k)
            except Exception as e:
                return self._json({"ok": False,
                                   "msg": f"{type(e).__name__}: {e}"}, 500)
            out = {
                "ok": True, "mode": mode, "tag": tag,
                "ms": round((time.perf_counter() - t0) * 1000),
                "hits": [{"sid": r["sid"], "score": r["_score"]} for r in rows],
            }
            # 向量比資料舊但在容許範圍內（只警告、照常查詢）時，把各幾筆告訴網頁，
            # 讓使用者看得到，而不是只有終端機有一行字。
            if mode in ("semantic", "hybrid"):
                st = index().vector_info(tag)
                if not st["exact"]:
                    out["stale"] = st
            return self._json(out)
        return super().do_GET()

    def _origin_ok(self):
        """POST 會真的做事（開始抓取、停止、寫檔），只收自己人送來的。
        瀏覽器對「別的網站」送來的簡單 POST 照樣會送到，只是看不到回應 ——
        所以光靠 CORS 擋不住，要在這裡看 Origin。沒有 Origin 的是命令列工具，放行。"""
        origin = self.headers.get("Origin")
        if not origin or origin in self.ALLOW_ORIGINS:
            return True
        # 區網其他電腦用 http://192.168.x.x:8765 打開的網頁，Origin 就是它自己。
        # Host 要是 IP 才算：用網域名稱指到區網位址（DNS rebinding）的不算。
        host = (self.headers.get("Host") or "").rsplit(":", 1)[0]
        try:
            ipaddress.ip_address(host)
        except ValueError:
            return False
        return origin == f"http://{self.headers.get('Host')}"

    def do_POST(self):
        path, _, qs = self.path.partition("?")
        q = urllib.parse.parse_qs(qs)
        if not self._origin_ok():
            return self._json({"ok": False, "msg": "來源不被允許"}, 403)
        if path == "/api/xlsx":
            # 網頁把「目前畫面上篩出來的那些列」送過來，這裡寫成 xlsx。
            # 瀏覽器自己做不出真正的 Excel 檔（要額外的函式庫），但後端
            # 本來就有 openpyxl，交給它最省事也最不會有編碼問題。
            try:
                n = int(self.headers.get("Content-Length") or 0)
                body = json.loads(self.rfile.read(n).decode("utf-8"))
                name = write_xlsx(body)
            except Exception as e:
                return self._json({"ok": False, "msg": f"{type(e).__name__}: {e}"}, 500)
            return self._json({"ok": True, "file": name})
        if path == "/api/stop":
            return self._json({"ok": JOB.stop()})
        if path == "/api/tool":
            key = q.get("key", [""])[0]
            if key not in TOOLS:
                return self._json({"ok": False, "msg": "不認得的工具"}, 400)
            label, args = TOOLS[key]
            ok, msg = JOB.start(label, args)
            return self._json({"ok": ok, "msg": msg}, 200 if ok else 409)
        if path == "/api/run":
            label, args = build_args(q)
            then = FOLLOW["sync"] if q.get("sync", [""])[0] == "1" else None
            ok, msg = JOB.start(label, args, then)
            return self._json({"ok": ok, "msg": msg}, 200 if ok else 409)
        return self.send_error(404)

    # 允許哪些網頁來呼叫這台本機工作台。
    # 瀏覽器禁止 https 網頁連 http 網站，但 127.0.0.1 是例外（本機被當成
    # 可信任來源），所以 GitHub Pages 上的那頁可以直接驅動這裡的爬蟲。
    # "null" 是在資料夾裡點兩下 HTML（file://）時瀏覽器送出的 Origin。
    # 不用 "*"：這台伺服器雖然只聽 127.0.0.1，但任何網站都能叫你的瀏覽器
    # 去打它，開放全部等於誰都能在你電腦上按下「開始抓取」。
    ALLOW_ORIGINS = {
        "null",
        "https://andy939.github.io",
        "http://127.0.0.1:8765", "http://localhost:8765",
    }

    def _cors(self):
        origin = self.headers.get("Origin")
        if origin in self.ALLOW_ORIGINS:
            self.send_header("Access-Control-Allow-Origin", origin)
            self.send_header("Vary", "Origin")

    def do_OPTIONS(self):
        # 預檢請求：瀏覽器在送 POST 之前會先問一次「我可以嗎」。
        # 這裡不要自己叫 _cors()，end_headers() 已經會加 —— 加兩次會變成
        # 「Allow-Origin: null,null」，瀏覽器看到多值就整個判定無效。
        self.send_response(204)
        self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")
        self.send_header("Access-Control-Max-Age", "86400")
        self.end_headers()

    def end_headers(self):
        # 資料檔每次抓完就變，不要讓瀏覽器拿舊的
        if self.path.endswith((".json", ".csv", ".html")):
            self.send_header("Cache-Control", "no-cache")
        self._cors()
        super().end_headers()

    def log_message(self, *a):
        pass                                # 不要把每個請求都印出來洗版


class LanServer(ThreadingHTTPServer):
    """區網模式：聽所有網路卡，但只接「跟這台同一個網段」來的連線。
    這台是 192.168.1.33，就只收 192.168.1.x（前三段一樣），外加本機 127.x；
    其他一律在連線時就斷掉，連網頁都看不到。
    路由器沒有做連接埠轉送的話，外面本來也連不進來，這是第二道。"""

    def __init__(self, addr, handler, net):
        self.net = net
        super().__init__(addr, handler)

    def verify_request(self, request, client_address):
        try:
            ip = ipaddress.ip_address(client_address[0])
        except ValueError:
            return False
        return ip.is_loopback or ip in self.net


def lan_ip():
    """這台電腦在區網裡的位址。連一下外面的位址（UDP 不會真的送出封包）
    看系統選哪張網卡，比 gethostbyname 準 —— 那個常拿到虛擬網卡的位址。"""
    try:
        with socket.socket(socket.AF_INET, socket.SOCK_DGRAM) as s:
            s.connect(("8.8.8.8", 80))
            return s.getsockname()[0]
    except OSError:
        return socket.gethostbyname(socket.gethostname())


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--lan", action="store_true",
                    help="讓同一個區網的其他電腦也能連（區網以外一律擋掉）")
    a = ap.parse_args()
    if not os.path.exists(os.path.join(DOCS, "index.html")):
        sys.exit(f"找不到 {DOCS}\\index.html")
    url = f"http://127.0.0.1:{PORT}"
    if a.lan:
        ip = lan_ip()
        net = ipaddress.ip_network(f"{ip}/24", strict=False)   # 前三段一樣的
        srv = LanServer(("0.0.0.0", PORT), Handler, net)
        print("【區網模式】同一個網路的其他電腦，用瀏覽器打開：")
        print(f"\n        http://{ip}:{PORT}\n")
        print(f"只接受 {str(net.network_address).rsplit('.', 1)[0]}.x 來的連線，其他一律擋掉。")
        print("第一次啟動時 Windows 會問要不要允許，請只勾「私人網路」。\n")
    else:
        srv = ThreadingHTTPServer(("127.0.0.1", PORT), Handler)
    warm_up()
    print(f"本機工作台：{url}")
    print("網頁上的控制台可以直接按下去抓資料，也可以中途停止。")
    print("關掉這個視窗（或按 Ctrl-C）就停止。\n")
    threading.Timer(0.8, lambda: webbrowser.open(url)).start()
    try:
        srv.serve_forever()
    except KeyboardInterrupt:
        print("\n停止。")


if __name__ == "__main__":
    main()
