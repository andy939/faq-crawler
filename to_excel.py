#!/usr/bin/env python3
"""
把 faq.db 的測量資料輸出成一個 Excel 檔（多工作表 + 圖表 + 名詞說明）。

    python3 to_excel.py                 # 全部
    python3 to_excel.py --n 500         # 只要每次 500 筆的那些執行
    python3 to_excel.py --out 分析.xlsx

需要 openpyxl：pip install openpyxl
"""

import sys
import argparse
import os
import sqlite3
import statistics as S
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Windows 的主控台預設是 cp950，印繁體中文會 UnicodeEncodeError。
if sys.platform == "win32":
    for _s in (sys.stdout, sys.stderr):
        try:
            _s.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass


HERE = os.path.dirname(os.path.abspath(__file__))
DB = os.path.join(HERE, "faq.db")

NAMES = {"A": "每筆重開連線", "B": "直接抓・循序", "C": "關掉 gzip",
         "D": "完整瀏覽器流程", "E": "併發 3", "F": "併發 8"}

HDR_FILL = PatternFill("solid", fgColor="1F4E5F")
HDR_FONT = Font(color="FFFFFF", bold=True, size=10)
THIN = Side(style="thin", color="D0D5DA")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(ws, ncol, row=1):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def autofit(ws, maxw=46):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        best = 0
        for c in col[:400]:
            v = c.value
            if v is None:
                continue
            # 中文字大約佔兩個字元寬
            t = str(v)
            best = max(best, sum(2 if ord(ch) > 0x2E80 else 1 for ch in t))
        ws.column_dimensions[letter].width = min(max(best + 2, 8), maxw)


def write(ws, header, rows, numfmt=None):
    ws.append(header)
    for r in rows:
        ws.append(r)
    style_header(ws, len(header))
    for col, fmt in (numfmt or {}).items():
        i = header.index(col) + 1
        for row in range(2, ws.max_row + 1):
            ws.cell(row=row, column=i).number_format = fmt
    autofit(ws)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--n", type=int, default=0)
    ap.add_argument("--out", default=os.path.join(HERE, "爬蟲分析.xlsx"))
    a = ap.parse_args()

    con = sqlite3.connect(DB)
    con.row_factory = sqlite3.Row
    smp = [dict(r) for r in con.execute(
        "SELECT * FROM samples ORDER BY run_id, seq")]
    runs = [dict(r) for r in con.execute("SELECT * FROM runs ORDER BY id")]
    con.close()
    if not smp:
        raise SystemExit("samples 是空的，先在 UI 跑幾輪")

    have = {r["run_id"] for r in smp}
    runs = [r for r in runs if r["id"] in have]
    if a.n:
        runs = [r for r in runs if r["done"] >= a.n * 0.9]
        ids = {r["id"] for r in runs}
        smp = [r for r in smp if r["run_id"] in ids]

    by_m = defaultdict(list)
    for r in smp:
        by_m[r["method"]].append(r)
    methods = sorted(by_m)
    runs_by_m = defaultdict(list)
    for r in runs:
        runs_by_m[r["method"]].append(r)

    wb = Workbook()

    # ---------- 說明 ----------
    ws = wb.active
    ws.title = "說明"
    gloss = [
        ("名詞", "意思"),
        ("實際耗時（牆鐘）",
         "從開始到結束，牆上的時鐘走了多久。就是你實際等待的時間。"
         "併發時多個請求同時進行，所以會小於「所有請求時間加總」。"),
        ("等伺服器",
         "送出請求 → 收到第一個位元組。包含建立連線（TCP/TLS）"
         "和伺服器組出頁面的時間。也叫 TTFB。"),
        ("傳輸", "收到第一個位元組 → 把整個頁面收完。純粹是資料流過網路的時間。"),
        ("單次總計", "等伺服器 + 傳輸。一次請求從頭到尾。"),
        ("抓下位元數", "真正流過網路的位元組（gzip 壓縮後）。"),
        ("網頁原始位元數", "解壓後的完整 HTML。關掉 gzip 時兩者相同。"),
        ("需要位元數",
         "發布時間＋更新時間＋標題＋發布機關＋內容 這 5 個欄位的 UTF-8 位元組。"
         "中文一字 3 bytes。"),
        ("起算秒 / 完成秒", "這一筆相對於該次執行開始的秒數，可用來還原併發的時間軸。"),
        ("gzip", "一種壓縮格式。伺服器壓過再送，收到端自動解開。這站約壓到 28%。"),
        ("keep_alive", "同一條連線重複使用，不用每次重新握手。"),
        ("", ""),
        ("抓法", "說明"),
    ] + [(m, NAMES.get(m, "")) for m in methods] + [
        ("", ""),
        ("工作表", "內容"),
        ("方法比較", "每種抓法的彙總數字"),
        ("逐筆明細", "每次執行 × 每一筆的原始測量，可直接做樞紐分析"),
        ("執行紀錄", "每次執行的彙總與設定"),
        ("圖表", "六張圖：每筆耗時、時間拆解、位元組對比、併發效益、"
                 "內容分組、方法 A 被限流的過程"),
        ("圖表資料", "上面那些圖的來源數字（預設隱藏，要看可以取消隱藏）"),
    ]
    for r in gloss:
        ws.append(list(r))
    for row in (1, len(gloss) - len(methods) - 6, len(gloss) - 4):
        pass
    style_header(ws, 2)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 90
    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=2).alignment = Alignment(wrap_text=True,
                                                         vertical="top")
        ws.cell(row=row, column=1).font = Font(bold=True)

    # ---------- 方法比較 ----------
    ws = wb.create_sheet("方法比較")
    hdr = ["抓法", "說明", "併發數", "keep_alive", "gzip", "筆數",
           "實際耗時秒", "每筆耗時ms", "等伺服器ms", "傳輸ms", "單次總計ms",
           "抓下位元數", "原始HTML位元數", "需要位元數", "壓縮率%", "有效率1/N",
           "req/s"]
    rows = []
    for m in methods:
        rs = by_m[m]
        ru = runs_by_m[m]
        wall = sum(x["elapsed"] for x in ru)
        n = len(rs)
        med = lambda k: S.median(x[k] for x in rs)
        mw, mr, mn = med("bytes_wire"), med("bytes_raw"), med("bytes_needed")
        rows.append([
            m, NAMES.get(m, ""), rs[0]["workers"], rs[0]["keep_alive"],
            rs[0]["gzip"], n, round(wall, 1), round(wall / n * 1000),
            round(med("t_wait") * 1000, 1), round(med("t_read") * 1000, 1),
            round(med("t_total") * 1000, 1), int(mw), int(mr), int(mn),
            round(100 * mw / mr, 1), round(mw / mn), round(n / wall, 1)])
    write(ws, hdr, rows, {"抓下位元數": "#,##0", "原始HTML位元數": "#,##0",
                          "需要位元數": "#,##0"})

    # ---------- 逐筆明細 ----------
    ws = wb.create_sheet("逐筆明細")
    cols = [("run_id", "run_id"), ("seq", "序"), ("method", "抓法"),
            ("workers", "併發數"), ("gzip", "gzip"), ("keep_alive", "keep_alive"),
            ("published", "發布時間"), ("updated", "更新時間"),
            ("title", "標題"), ("dept", "發布機關"), ("chars", "內容字數"),
            ("bytes_wire", "抓下位元數"), ("bytes_raw", "網頁原始位元數"),
            ("bytes_needed", "需要位元數"), ("t_wait", "等伺服器ms"),
            ("t_read", "傳輸ms"), ("t_total", "單次總ms"),
            ("t_start", "起算秒"), ("t_end", "完成秒"),
            ("worker", "執行緒"), ("sid", "sid")]
    ms = {"t_wait", "t_read", "t_total"}
    sec = {"t_start", "t_end"}

    def conv(k, v):
        if v is None:
            return ""
        if k in ms:
            return round(v * 1000, 1)
        if k in sec:
            return round(v, 2)
        return v

    write(ws, [zh for _, zh in cols],
          [[conv(en, r[en]) for en, _ in cols] for r in smp],
          {"抓下位元數": "#,##0", "網頁原始位元數": "#,##0", "需要位元數": "#,##0"})
    ws.auto_filter.ref = ws.dimensions

    # ---------- 執行紀錄 ----------
    ws = wb.create_sheet("執行紀錄")
    rcols = ["id", "ts", "method", "label", "note", "n", "done", "elapsed",
             "reqs", "wire", "raw", "kept", "workers", "gzip", "keep_alive",
             "warm_session", "hit_counter", "lat_min", "lat_p50", "lat_p95",
             "lat_max", "errors"]
    zh = ["編號", "時間", "抓法", "說明", "範圍", "要求筆數", "完成筆數",
          "實際耗時秒", "請求數", "抓下位元數", "原始位元數", "需要位元數",
          "併發數", "gzip", "keep_alive", "帶cookie", "點閱回報",
          "最快秒", "中位秒", "p95秒", "最慢秒", "錯誤數"]
    write(ws, zh, [[round(r[c], 3) if isinstance(r[c], float) else r[c]
                    for c in rcols] for r in runs],
          {"抓下位元數": "#,##0", "原始位元數": "#,##0", "需要位元數": "#,##0"})

    # ---------- 圖表 ----------
    # 圖表的來源資料放在另一張（隱藏的）工作表，圖表頁只留圖，才不會一堆
    # 裸露的數字看起來像沒做完。
    dat = wb.create_sheet("圖表資料")
    ws = wb.create_sheet("圖表")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3

    dcol = [1]      # 圖表資料頁目前用到第幾欄
    anchor = [2]    # 圖表頁目前放到第幾列

    def block(header, rows):
        """把一塊資料寫進圖表資料頁，回傳 (起欄, 起列, 迄列)"""
        c0 = dcol[0]
        for j, h in enumerate(header):
            dat.cell(row=1, column=c0 + j, value=h)
        for i, r in enumerate(rows):
            for j, v in enumerate(r):
                dat.cell(row=2 + i, column=c0 + j, value=v)
        dcol[0] += len(header) + 1
        return c0, 1, 1 + len(rows)

    def place(chart, title, note=""):
        r = anchor[0]
        ws.cell(row=r, column=2, value=title).font = Font(bold=True, size=13)
        if note:
            ws.cell(row=r + 1, column=2, value=note).font = Font(
                size=10, color="6B7280")
        chart.height, chart.width = 9.5, 22
        ws.add_chart(chart, f"B{r + 2}")
        anchor[0] = r + 23        # 9.5cm 約 19 列，留 4 列間距

    mlabel = [f"{m}　{NAMES.get(m, '')}" for m in methods]

    def med(m, k):
        return S.median(x[k] for x in by_m[m])

    def wall(m):
        return sum(x["elapsed"] for x in runs_by_m[m])

    # 1 每筆耗時
    c0, r0, r1 = block(["抓法", "每筆耗時ms"],
                       [[mlabel[i], round(wall(m) / len(by_m[m]) * 1000)]
                        for i, m in enumerate(methods)])
    ch = BarChart(); ch.type = "col"
    ch.y_axis.title = "毫秒"; ch.legend = None
    ch.add_data(Reference(dat, min_col=c0 + 1, min_row=r0, max_row=r1),
                titles_from_data=True)
    ch.set_categories(Reference(dat, min_col=c0, min_row=r0 + 1, max_row=r1))
    place(ch, "1　每筆耗時（越低越好）",
          "實際耗時 ÷ 筆數。A 因為中途被限流所以特別高。")

    # 2 時間拆解
    c0, r0, r1 = block(["抓法", "等伺服器ms", "傳輸ms"],
                       [[mlabel[i], round(med(m, "t_wait") * 1000, 1),
                         round(med(m, "t_read") * 1000, 1)]
                        for i, m in enumerate(methods)])
    ch = BarChart(); ch.type = "col"; ch.grouping = "stacked"; ch.overlap = 100
    ch.y_axis.title = "毫秒（單次請求中位數）"
    ch.add_data(Reference(dat, min_col=c0 + 1, max_col=c0 + 2,
                          min_row=r0, max_row=r1), titles_from_data=True)
    ch.set_categories(Reference(dat, min_col=c0, min_row=r0 + 1, max_row=r1))
    place(ch, "2　一次請求的時間拆解",
          "C 慢在「傳輸」（沒壓縮）；A 慢在「等伺服器」（被限流）。")

    # 3 位元組
    c0, r0, r1 = block(["抓法", "抓下位元數", "需要位元數"],
                       [[mlabel[i], int(med(m, "bytes_wire")),
                         int(med(m, "bytes_needed"))]
                        for i, m in enumerate(methods)])
    ch = BarChart(); ch.type = "col"
    ch.y_axis.title = "位元組（中位數）"
    ch.add_data(Reference(dat, min_col=c0 + 1, max_col=c0 + 2,
                          min_row=r0, max_row=r1), titles_from_data=True)
    ch.set_categories(Reference(dat, min_col=c0, min_row=r0 + 1, max_row=r1))
    place(ch, "3　抓下的量 vs 真正需要的量",
          "「需要」那條幾乎貼在零 —— 500 byte 對 55,000 byte。")

    # 4 併發效益
    c0, r0, r1 = block(["抓法", "請求時間總和秒", "實際耗時秒"],
                       [[mlabel[i], round(sum(x["t_total"] for x in by_m[m]), 1),
                         round(wall(m), 1)] for i, m in enumerate(methods)])
    ch = BarChart(); ch.type = "col"
    ch.y_axis.title = "秒"
    ch.add_data(Reference(dat, min_col=c0 + 1, max_col=c0 + 2,
                          min_row=r0, max_row=r1), titles_from_data=True)
    ch.set_categories(Reference(dat, min_col=c0, min_row=r0 + 1, max_row=r1))
    place(ch, "4　併發省了多少",
          "兩條落差越大，代表併行度越高。E/F 的實際耗時遠低於總和。")

    # 5 有效率 vs 內容大小
    bs = sorted(by_m[methods[0]], key=lambda r: r["bytes_needed"])
    k = max(len(bs) // 5, 1)
    labs = ["最小 20%", "20~40%", "40~60%", "60~80%", "最大 20%"]
    rows5 = []
    for i, lab in enumerate(labs):
        g = bs[i * k:(i + 1) * k] if i < 4 else bs[4 * k:]
        if g:
            rows5.append([lab, int(S.median(r["bytes_needed"] for r in g)),
                          int(S.median(r["bytes_wire"] for r in g))])
    c0, r0, r1 = block(["分組", "需要位元數", "抓下位元數"], rows5)
    ch = BarChart(); ch.type = "col"
    ch.y_axis.title = "位元組"
    ch.add_data(Reference(dat, min_col=c0 + 1, max_col=c0 + 2,
                          min_row=r0, max_row=r1), titles_from_data=True)
    ch.set_categories(Reference(dat, min_col=c0, min_row=r0 + 1, max_row=r1))
    place(ch, "5　內容大小分組：抓下的量幾乎不變",
          "內容從 243 到 1,418 byte，抓下的量只從 54.2K 變到 56.4K。")

    # 6 A 的限流曲線
    if "A" in by_m and len(by_m["A"]) > 50:
        rs = sorted(by_m["A"], key=lambda r: r["t_start"])
        c0, r0, r1 = block(["序", "等伺服器ms"],
                           [[i + 1, round(r["t_wait"] * 1000, 1)]
                            for i, r in enumerate(rs)])
        lc = LineChart()
        lc.y_axis.title = "等伺服器（毫秒）"; lc.x_axis.title = "第幾筆"
        lc.legend = None
        lc.add_data(Reference(dat, min_col=c0 + 1, min_row=r0, max_row=r1),
                    titles_from_data=True)
        lc.set_categories(Reference(dat, min_col=c0, min_row=r0 + 1,
                                    max_row=r1))
        for s_ in lc.series:
            s_.smooth = False
            s_.graphicalProperties.line.width = 12000
        place(lc, "6　方法 A 被限流的過程",
              "前 89 筆平貼在 180ms，第 90 筆起垂直跳到約 2,200ms 就固定不動 "
              "—— 階梯形狀是限流，不是伺服器忙。")

    dat.sheet_state = "hidden"

    wb.save(a.out)
    print(f"已產生 {a.out}")
    print(f"  工作表：{'、'.join(wb.sheetnames)}")
    print(f"  逐筆明細 {len(smp)} 列，執行紀錄 {len(runs)} 列")


if __name__ == "__main__":
    main()
